# -*- coding: utf-8 -*-
"""Monta o workbook Excel (Leia-me, Resumo, PF, PJ, Listas) a partir dos lançamentos."""
from datetime import date, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .pipeline import DEFAULT_CATEGORIES

ARIAL = "Arial"
NAVY, BLUE, LGREY, HGREY = "1F3864", "2E5496", "F2F2F2", "D9E1F2"
GREEN, RED = "2E7D32", "C0392B"
MONEY = "R$ #,##0.00;[RED]-R$ #,##0.00"
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(_thin, _thin, _thin, _thin)


def _f(sz=10, b=False, c="000000"):
    return Font(name=ARIAL, size=sz, bold=b, color=c)


def _fill(c):
    return PatternFill("solid", fgColor=c)


COLS = [("Data", 13), ("Fonte", 22), ("Descrição", 40), ("Categoria", 30),
        ("Tipo", 34), ("Entrada (R$)", 14), ("Saída (R$)", 14), ("Mês", 9), ("Obs", 26)]


def _ledger(wb, name, escopo, tx):
    w = wb.create_sheet(name)
    w.sheet_view.showGridLines = False
    for j, (h, wd) in enumerate(COLS, start=1):
        c = w.cell(1, j, h)
        c.font = _f(10, True, "FFFFFF"); c.fill = _fill(BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        w.column_dimensions[get_column_letter(j)].width = wd
    w.row_dimensions[1].height = 30
    rows = [t for t in tx if t["escopo"] == escopo]
    ri = 2
    for t in rows:
        d = t["data"]
        if isinstance(d, datetime):
            d = d.date()
        w.cell(ri, 1, d).number_format = "dd/mm/yyyy"
        w.cell(ri, 2, t["fonte"])
        w.cell(ri, 3, t["descricao"])
        w.cell(ri, 4, t["categoria"])
        w.cell(ri, 5, t["tipo"])
        ce = w.cell(ri, 6, t["entrada"] or None); ce.number_format = MONEY; ce.font = _f(10, c=GREEN)
        cs = w.cell(ri, 7, t["saida"] or None); cs.number_format = MONEY; cs.font = _f(10, c=RED)
        w.cell(ri, 8, f"{d.year}-{d.month:02d}" if d else "")
        w.cell(ri, 9, t["obs"])
        for j in range(1, 10):
            cell = w.cell(ri, j)
            cell.border = BORDER
            if cell.font.name != ARIAL:
                cell.font = _f(10)
        ri += 1
    last = ri - 1
    tc = w.cell(ri, 5, "TOTAL"); tc.font = _f(10, True); tc.alignment = Alignment(horizontal="right")
    for col, ltr in [(6, "F"), (7, "G")]:
        cc = w.cell(ri, col, f"=SUM({ltr}2:{ltr}{last})")
        cc.number_format = MONEY; cc.font = _f(10, True); cc.fill = _fill(HGREY)
    w.freeze_panes = "A2"
    w.auto_filter.ref = f"A1:I{last}"
    return last


def build_workbook(tx, out_path, extra_categories=None):
    wb = openpyxl.Workbook()

    # ---- Leia-me ----
    ws = wb.active; ws.title = "Leia-me"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 110
    L = [("Planilha de Despesas e Movimentações", 15, True, NAVY),
         ("Gerada automaticamente pelo pipeline (contabil).", 9, False, "000000"),
         ("", 10, False, "000000"),
         ("Abas PF e PJ: todos os lançamentos por titularidade da conta (PF = Iuri, PJ = Youup).", 10, False, "000000"),
         ("Resumo: receitas/despesas, despesas por categoria e movimentação por mês (fórmulas).", 10, False, "000000"),
         ("Listas: categorias-padrão usadas no menu suspenso da coluna Categoria.", 10, False, "000000"),
         ("", 10, False, "000000"),
         ("Coluna Tipo:", 11, True, BLUE),
         ("Despesa/Receita = somam nos totais. Estorno = crédito que reduz despesa.", 10, False, "000000"),
         ("Transferência entre contas próprias / Pagamento de fatura / Liquidação = memo (não somam).", 10, False, "000000"),
         ("", 10, False, "000000"),
         ("Categorias são sugestões automáticas; revise as marcadas 'A classificar'.", 10, False, RED)]
    for i, (txt, sz, b, c) in enumerate(L, start=2):
        cell = ws.cell(i, 2, txt); cell.font = _f(sz, b, c)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- PF / PJ ----
    pfl = _ledger(wb, "PF", "PF", tx)
    pjl = _ledger(wb, "PJ", "PJ", tx)

    # ---- Listas ----
    cats = sorted(set(DEFAULT_CATEGORIES) | set(extra_categories or []) |
                  {t["categoria"] for t in tx})
    wl = wb.create_sheet("Listas")
    wl.cell(1, 1, "Categorias (edite/adicione abaixo)").font = _f(11, True, NAVY)
    wl.cell(1, 3, "Adicione novas categorias em células vazias da coluna A: elas entram no "
                  "menu suspenso e ganham linha no Resumo.").font = _f(9, c="7030A0")
    for i, c in enumerate(cats, start=2):
        wl.cell(i, 1, c).font = _f(10)
    wl.column_dimensions["A"].width = 46
    LB = 100  # faixa do dropdown / espelho no Resumo

    # dropdown na coluna Categoria (PF e PJ)
    lr = f"Listas!$A$2:$A${LB}"
    for sh, last in (("PF", pfl), ("PJ", pjl)):
        if last < 2:
            continue  # aba sem lançamentos: não há intervalo válido para o dropdown
        w = wb[sh]
        dv = DataValidation(type="list", formula1=f"={lr}", allow_blank=True, showDropDown=False)
        w.add_data_validation(dv); dv.add(f"D2:D{last}")

    # ---- Resumo ----
    ws = wb.create_sheet("Resumo"); ws.sheet_view.showGridLines = False
    for col, wd in [("A", 2), ("B", 42), ("C", 16), ("D", 16), ("E", 2), ("F", 42), ("G", 16), ("H", 16)]:
        ws.column_dimensions[col].width = wd

    def hdr(r, c, t):
        x = ws.cell(r, c, t); x.font = _f(10, True, "FFFFFF"); x.fill = _fill(BLUE)
        x.alignment = Alignment(horizontal="center"); x.border = BORDER

    ws.cell(2, 2, "Resumo geral").font = _f(14, True, NAVY)
    hdr(4, 2, "Visão geral"); hdr(4, 3, "PF (Iuri)"); hdr(4, 4, "PJ (Youup)")
    r = 5
    for lab, tp, col in [("Receitas (entradas)", "Receita", "F"),
                         ("Despesas (saídas)", "Despesa", "G"),
                         ("Estornos/Devoluções (crédito)", "Estorno/Devolução", "F")]:
        ws.cell(r, 2, lab).font = _f(10); ws.cell(r, 2).border = BORDER
        for cc, sh, last in [(3, "PF", pfl), (4, "PJ", pjl)]:
            f = f'=SUMIFS({sh}!${col}$2:${col}${last},{sh}!$E$2:$E${last},"{tp}")'
            x = ws.cell(r, cc, f); x.number_format = MONEY; x.font = _f(10); x.border = BORDER
        r += 1
    ws.cell(r, 2, "Resultado (Receitas − Despesas)").font = _f(10, True)
    ws.cell(r, 2).fill = _fill(HGREY); ws.cell(r, 2).border = BORDER
    for cc in (3, 4):
        LC = get_column_letter(cc)
        x = ws.cell(r, cc, f"={LC}5-{LC}6"); x.number_format = MONEY; x.font = _f(10, True)
        x.fill = _fill(HGREY); x.border = BORDER
    r += 2
    ws.cell(r, 2, "Memorando (não somam)").font = _f(11, True, BLUE); r += 1
    hdr(r, 2, "Item"); hdr(r, 3, "PF"); hdr(r, 4, "PJ"); r += 1
    for lab, tp, col in [("Transferências entre contas próprias", "Transferência entre contas próprias", "G"),
                         ("Pagamentos de fatura de cartão", "Pagamento de fatura de cartão", "F"),
                         ("Liquidações internas (adquirente)", "Liquidação/Interno (adquirente)", "G")]:
        ws.cell(r, 2, lab).font = _f(10); ws.cell(r, 2).border = BORDER
        for cc, sh, last in [(3, "PF", pfl), (4, "PJ", pjl)]:
            f = f'=SUMIFS({sh}!${col}$2:${col}${last},{sh}!$E$2:$E${last},"{tp}")'
            x = ws.cell(r, cc, f); x.number_format = MONEY; x.font = _f(10); x.border = BORDER
        r += 1
    r += 1
    ws.cell(r, 2, "Movimentação por mês (despesas)").font = _f(12, True, NAVY); r += 1
    hdr(r, 2, "Mês"); hdr(r, 3, "PF"); hdr(r, 4, "PJ"); r += 1
    ms = r
    nomes = {1: "Jan", 2: "Fev", 3: "Mar", 4: "Abr", 5: "Mai", 6: "Jun", 7: "Jul", 8: "Ago",
             9: "Set", 10: "Out", 11: "Nov", 12: "Dez"}
    for mth in range(1, 13):
        key = f"2026-{mth:02d}"
        ws.cell(r, 2, f"{nomes[mth]}/26").font = _f(10); ws.cell(r, 2).border = BORDER
        for cc, sh, last in [(3, "PF", pfl), (4, "PJ", pjl)]:
            f = f'=SUMIFS({sh}!$G$2:$G${last},{sh}!$E$2:$E${last},"Despesa",{sh}!$H$2:$H${last},"{key}")'
            x = ws.cell(r, cc, f); x.number_format = MONEY; x.font = _f(10); x.border = BORDER
        r += 1
    ws.cell(r, 2, "Total").font = _f(10, True); ws.cell(r, 2).fill = _fill(HGREY); ws.cell(r, 2).border = BORDER
    for cc in (3, 4):
        LC = get_column_letter(cc)
        x = ws.cell(r, cc, f"=SUM({LC}{ms}:{LC}{r-1})"); x.number_format = MONEY
        x.font = _f(10, True); x.fill = _fill(HGREY); x.border = BORDER

    # despesas por categoria (espelha Listas A2:A100)
    ws.cell(4, 6, "Despesas por categoria (espelha a aba Listas)").font = _f(12, True, NAVY)
    hdr(5, 6, "Categoria"); hdr(5, 7, "PF (R$)"); hdr(5, 8, "PJ (R$)")
    rr = 6
    for k in range(2, LB + 1):
        ws.cell(rr, 6, f'=IF(Listas!$A{k}="","",Listas!$A{k})').font = _f(10)
        ws.cell(rr, 6).border = BORDER
        gF = f'=IF($F{rr}="","",SUMIFS(PF!$G$2:$G${pfl},PF!$E$2:$E${pfl},"Despesa",PF!$D$2:$D${pfl},$F{rr}))'
        hF = f'=IF($F{rr}="","",SUMIFS(PJ!$G$2:$G${pjl},PJ!$E$2:$E${pjl},"Despesa",PJ!$D$2:$D${pjl},$F{rr}))'
        x = ws.cell(rr, 7, gF); x.number_format = MONEY; x.font = _f(10); x.border = BORDER
        y = ws.cell(rr, 8, hF); y.number_format = MONEY; y.font = _f(10); y.border = BORDER
        rr += 1
    ws.cell(rr, 6, "Total despesas").font = _f(10, True); ws.cell(rr, 6).fill = _fill(HGREY)
    ws.cell(rr, 6).border = BORDER
    for c in (7, 8):
        LC = get_column_letter(c)
        x = ws.cell(rr, c, f"=SUM({LC}6:{LC}{rr-1})"); x.number_format = MONEY
        x.font = _f(10, True); x.fill = _fill(HGREY); x.border = BORDER

    # ordem das abas
    order = {"Leia-me": 0, "Resumo": 1, "PF": 2, "PJ": 3, "Listas": 4}
    wb._sheets.sort(key=lambda s: order.get(s.title, 9))
    wb.save(out_path)
    return out_path
